import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))
import import_inventory as importer


class ParsingTests(unittest.TestCase):
    def test_identifier_keeps_leading_zeroes_and_nfkc(self):
        self.assertEqual(importer.normalize_identifier(" ０0123 "), "00123")

    def test_quantity_preserves_fraction_and_flags_bad_values(self):
        self.assertEqual(importer.parse_quantity("1.25个"), (1.25, None))
        self.assertEqual(importer.parse_quantity("0"), (0.0, "zero_quantity"))
        self.assertEqual(importer.parse_quantity("-2"), (-2.0, "negative_quantity"))
        self.assertEqual(importer.parse_quantity("两箱"), (None, "invalid_quantity"))

    def test_date_quality(self):
        self.assertEqual(importer.parse_date("2024.01.02"), ("2024-01-02", None))
        self.assertEqual(importer.parse_date("2323.01.02")[1], "out_of_range_date")
        self.assertEqual(importer.parse_date(None), (None, "missing_date"))

    def test_transfer_classification_is_conservative(self):
        self.assertEqual(importer.classify_transfer("备用仓库"), ("TRANSFER", "keyword_transfer"))
        self.assertEqual(importer.classify_transfer("已消耗"), ("ISSUE_CONSUMPTION", "keyword_terminal_outflow"))
        self.assertEqual(importer.classify_transfer("位置待定"), ("REVIEW", "no_safe_transfer_semantics"))

    def test_sensitive_values_are_not_emitted(self):
        self.assertEqual(importer.redact_value("contact", "联系方式"), "[REDACTED]")
        self.assertEqual(importer.redact_value("金额已付款", "备注"), "[REDACTED_FINANCIAL_OR_PERSONAL_NOTE]")
        self.assertEqual(importer.redact_value("通用备注", "备注"), "通用备注")

    def _mapping(self):
        return {
            "version": 1,
            "sheets": [
                {
                    "name": "Items",
                    "role": "products",
                    "header_row": 1,
                    "columns": {
                        "identifier": "SKU",
                        "name": "Product",
                        "uom": "Unit",
                        "opening_quantity": "Opening",
                        "existing_quantity": "On hand",
                    },
                },
                {
                    "name": "Inbound log",
                    "role": "movement",
                    "movement_type": "RECEIPT",
                    "columns": {
                        "identifier": "SKU",
                        "name": "Product",
                        "date": "Date",
                        "quantity": "Quantity",
                        "uom": "Unit",
                        "serial": "Serial",
                    },
                },
                {
                    "name": "Relocations",
                    "role": "movement",
                    "movement_type": "TRANSFER",
                    "columns": {
                        "identifier": "SKU",
                        "name": "Product",
                        "date": "Date",
                        "quantity": "Quantity",
                        "uom": "Unit",
                        "note": "Note",
                        "destination": "Destination",
                    },
                },
                {
                    "name": "Equipment",
                    "role": "assets",
                    "columns": {
                        "name": "Asset",
                        "serial": "Serial",
                        "component_serial": "Component serial",
                        "notes": "Notes",
                    },
                },
                {
                    "name": "References",
                    "role": "stage_only",
                    "columns": {
                        "reference": "Reference",
                        "contact": "Contact phone",
                        "amount": "Amount",
                        "public_note": "Description",
                    },
                },
            ],
        }

    def _workbook(self, path: Path):
        workbook = Workbook()
        products = workbook.active
        products.title = "Items"
        products.append(["SKU", "Product", "Unit", "Opening", "On hand", "Internal comment"])
        products.append(["0001", "Sample item", "EA", 2, 10, "private source detail"])
        products.append(["0002", "Another item", "BOX", None, 4, "not exported"])

        receipts = workbook.create_sheet("Inbound log")
        receipts.append(["SKU", "Product", "Date", "Quantity", "Unit", "Serial", "Contact phone"])
        receipts.append(["0001", "Sample item", "2024-01-02", "1.25", "EA", "SN-ABC123456", "13812345678"])
        receipts.append(["0001", "Sample item", "2323-01-02", "0", "EA", "", "13812345678"])

        moves = workbook.create_sheet("Relocations")
        moves.append(["SKU", "Product", "Date", "Quantity", "Unit", "Note", "Destination"])
        moves.append(["0001", "Sample item", "2024-02-03", 1, "EA", "移至备用区", "Shelf B"])
        moves.append(["0001", "Sample item", "2024-02-04", 1, "EA", "位置待定", ""])

        assets = workbook.create_sheet("Equipment")
        assets.append(["Asset", "Serial", "Component serial", "Notes"])
        assets.append(["Sample equipment", "ASSET-123456", "PART-654321", "Synthetic fixture"])

        refs = workbook.create_sheet("References")
        refs.append(["Reference", "Contact phone", "Amount", "Description"])
        refs.append(["REF-01", "13812345678", 125.50, 13812345678])

        unknown = workbook.create_sheet("Unmapped")
        unknown.append(["Personal note"])
        unknown.append(["private@example.invalid"])
        workbook.save(path)
        workbook.close()

    def test_generic_mapping_imports_only_selected_fields(self):
        with tempfile.TemporaryDirectory(prefix="mini-erp-import-test-") as temp:
            root = Path(temp)
            workbook_path = root / "synthetic.xlsx"
            mapping_path = root / "mapping.json"
            self._workbook(workbook_path)
            mapping_path.write_text(json.dumps(self._mapping()), encoding="utf-8")

            report = importer.build_report(workbook_path, "", mapping_path)

            self.assertEqual(len(report["product_observations"]), 2)
            self.assertEqual(report["product_observations"][0]["identifier_normalized"], "0001")
            self.assertEqual(report["product_observations"][0]["source"]["existing_quantity"], 10)
            self.assertEqual(len(report["movement_candidates"]), 4)
            self.assertEqual(report["movement_candidates"][0]["movement_type_candidate"], "RECEIPT")
            self.assertEqual(report["movement_candidates"][0]["serial_candidates"], ["SN-ABC123456"])
            self.assertIn("out_of_range_date", report["movement_candidates"][1]["data_quality_issues"])
            self.assertIn("zero_quantity", report["movement_candidates"][1]["data_quality_issues"])
            self.assertEqual(report["movement_candidates"][2]["movement_type_candidate"], "TRANSFER")
            self.assertEqual(report["movement_candidates"][3]["movement_type_candidate"], "REVIEW")
            self.assertEqual(report["asset_observations"][0]["serial_candidates"], ["ASSET-123456", "PART-654321"])
            self.assertEqual(report["stage_only"][0]["source"], {"reference": "REF-01", "public_note": "[REDACTED_PHONE]"})
            self.assertEqual(report["summary"]["unmapped_sheets"], [{"name": "Unmapped", "row_count": 2}])
            serialized = json.dumps(report, ensure_ascii=False)
            self.assertNotIn("private source detail", serialized)
            self.assertNotIn("13812345678", serialized)
            self.assertNotIn("125.5", serialized)
            self.assertNotIn("private@example.invalid", serialized)

    def test_configured_one_based_header_row(self):
        with tempfile.TemporaryDirectory(prefix="mini-erp-import-test-") as temp:
            root = Path(temp)
            workbook_path = root / "synthetic.xlsx"
            mapping_path = root / "mapping.json"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Catalog"
            sheet.append(["Generated sample catalog"])
            sheet.append(["Code", "Label"])
            sheet.append(["0007", "Mapped item"])
            workbook.save(workbook_path)
            workbook.close()
            mapping_path.write_text(json.dumps({
                "version": 1,
                "sheets": [{
                    "name": "Catalog",
                    "role": "products",
                    "header_row": 2,
                    "columns": {"identifier": "Code", "name": "Label"},
                }],
            }), encoding="utf-8")

            report = importer.build_report(workbook_path, "", mapping_path)

            self.assertEqual(report["product_observations"][0]["identifier_normalized"], "0007")
            self.assertEqual(report["product_observations"][0]["source_row_number"], 3)

    def test_missing_mapped_header_fails_clearly(self):
        with tempfile.TemporaryDirectory(prefix="mini-erp-import-test-") as temp:
            root = Path(temp)
            workbook_path = root / "synthetic.xlsx"
            mapping_path = root / "mapping.json"
            self._workbook(workbook_path)
            mapping = self._mapping()
            mapping["sheets"][0]["columns"]["name"] = "Missing header"
            mapping_path.write_text(json.dumps(mapping), encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "缺少映射列"):
                importer.build_report(workbook_path, "", mapping_path)

    def test_mapping_requires_version_one_and_positive_header_row(self):
        with tempfile.TemporaryDirectory(prefix="mini-erp-import-test-") as temp:
            mapping_path = Path(temp) / "mapping.json"
            mapping = self._mapping()
            mapping["version"] = 2
            mapping_path.write_text(json.dumps(mapping), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "version 必须为 1"):
                importer.load_mapping(mapping_path)

            mapping["version"] = 1
            mapping["sheets"][0]["header_row"] = 0
            mapping_path.write_text(json.dumps(mapping), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "header_row 必须是正整数"):
                importer.load_mapping(mapping_path)


if __name__ == "__main__":
    unittest.main()
