"""批次3 管道与完整性测试：歧义日期、多行聚合、seed 审计判重、party 外键、过账人外键。"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))

from tests.support.testdb import DbTestCase

MAPPING = {
    "version": 1,
    "sheets": [
        {"name": "Items", "role": "products", "header_row": 1,
         "columns": {"identifier": "SKU", "name": "Product", "uom": "Unit",
                     "opening_quantity": "Opening", "existing_quantity": "On hand"}},
        {"name": "Inbound log", "role": "movement", "movement_type": "RECEIPT", "header_row": 1,
         "columns": {"identifier": "SKU", "name": "Product", "date": "Date",
                     "quantity": "Quantity", "uom": "Unit"}},
    ],
}


def _mapping_file(root: Path) -> Path:
    path = root / "mapping.json"
    path.write_text(json.dumps(MAPPING), encoding="utf-8")
    return path


class AmbiguousDateTests(unittest.TestCase):
    def test_ambiguous_slash_date_flagged_for_review(self) -> None:
        import import_inventory as importer
        value, issue = importer.parse_date("05/03/2025")
        self.assertEqual(issue, "ambiguous_date", "美式/欧式两解的日期必须转人工复核")
        self.assertIsNotNone(value)

    def test_unambiguous_slash_date_ok(self) -> None:
        import import_inventory as importer
        self.assertIsNone(importer.parse_date("25/03/2025")[1])
        self.assertIsNone(importer.parse_date("2024-01-02")[1])

    def test_dry_run_report_includes_product_quantity_issues(self) -> None:
        import import_inventory as importer
        with tempfile.TemporaryDirectory(prefix="p2-import-") as temp:
            root = Path(temp)
            workbook_path = root / "w.xlsx"
            wb = Workbook()
            sheet = wb.active
            sheet.title = "Items"
            sheet.append(["SKU", "Product", "Unit", "Opening", "On hand"])
            sheet.append(["0001", "数量问题货品", "EA", "两箱", -5])
            wb.create_sheet("Inbound log").append(["SKU", "Product", "Date", "Quantity", "Unit"])
            wb.save(workbook_path)
            wb.close()
            report = importer.build_report(workbook_path, "", _mapping_file(root))
            issues = report["product_observations"][0].get("data_quality_issues", [])
            self.assertIn("invalid_quantity", issues, "干运行报告应包含产品数量问题，供审核后再导入")
            self.assertIn("negative_quantity", issues)


class SeedPipelineTests(DbTestCase):
    def _workbook(self, root: Path, rows: list[tuple], movements: list[tuple]) -> Path:
        workbook = Workbook()
        products = workbook.active
        products.title = "Items"
        products.append(["SKU", "Product", "Unit", "Opening", "On hand"])
        for row in rows:
            products.append(list(row))
        receipts = workbook.create_sheet("Inbound log")
        receipts.append(["SKU", "Product", "Date", "Quantity", "Unit"])
        for row in movements:
            receipts.append(list(row))
        path = root / "w.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def _run(self, workbook_path: Path, mapping_path: Path) -> dict:
        import seed_inventory
        args = argparse.Namespace(input=str(workbook_path), mapping=str(mapping_path), password="",
                                  apply=True, post_opening=False, cutover_date=None)
        return seed_inventory.load(args)

    def _fetch_one(sql, params=()):  # noqa: N805
        from api.db import connection, fetch_one
        with connection() as conn:
            return fetch_one(conn, sql, params)

    def test_multi_row_product_snapshots_not_dropped(self) -> None:
        """同键多行（分批表达同一货品）的快照都要保留，不得只记第一行。"""
        import seed_inventory
        with tempfile.TemporaryDirectory(prefix="p2-seed-") as temp:
            root = Path(temp)
            workbook = self._workbook(root, [("7001", "多行货品", "EA", None, 10), ("7001", "多行货品", "EA", None, 4)], [])
            result = self._run(workbook, _mapping_file(root))
            from api.db import connection, fetch_all
            with connection() as conn:
                snaps = fetch_all(conn, """SELECT reported_quantity FROM inventory_snapshot s
                                             JOIN product p ON p.product_id=s.product_id
                                            WHERE p.display_name='多行货品' ORDER BY s.inventory_snapshot_id""")
            values = sorted(int(s["reported_quantity"]) for s in snaps)
            self.assertEqual(values, [4, 10], "同一货品的多行快照都应入库")

    def test_rerun_same_file_does_not_bloat_append_only_audit(self) -> None:
        with tempfile.TemporaryDirectory(prefix="p2-seed-") as temp:
            root = Path(temp)
            workbook = self._workbook(root, [("7002", "重跑货品", "EA", None, 3)],
                                      [("7002", "重跑货品", "2024-01-02", 2, "EA")])
            mapping = _mapping_file(root)
            self._run(workbook, mapping)
            from api.db import connection, fetch_one
            with connection() as conn:
                before = fetch_one(conn, "SELECT count(*) AS n FROM audit_event WHERE action IN ('SEED_SNAPSHOT','SEED_MOVEMENT','SEED_OPENING')")
            self._run(workbook, mapping)  # 同一文件（同 SHA）重跑
            with connection() as conn:
                after = fetch_one(conn, "SELECT count(*) AS n FROM audit_event WHERE action IN ('SEED_SNAPSHOT','SEED_MOVEMENT','SEED_OPENING')")
            self.assertEqual(before["n"], after["n"], "重复执行不应向 append-only 审计表灌重复行")

    def test_posted_by_user_id_recorded_for_api_movements(self) -> None:
        from tests.support.api_client import api_for
        admin = api_for("admin")
        warehouse = api_for("warehouse")
        loc = admin.post("/api/locations", json={"code": "PBULOC", "name": "过账人库位"}).json()["location_id"]
        ea = next(u for u in admin.get("/api/uoms").json() if u["code"] == "EA")
        product_id = admin.post("/api/products", json={"display_name": "过账人货品", "default_uom_id": ea["uom_id"], "source_uom_raw": "个"}).json()["product_id"]
        supplier_id = warehouse.post("/api/suppliers", json={"name": "过账人供应商"}).json()["supplier_id"]
        doc = warehouse.post("/api/documents", json={
            "doc_type": "PURCHASE_RECEIPT", "party_id": supplier_id,
            "lines": [{"product_id": product_id, "quantity": 2, "destination_location_id": loc}],
        }).json()
        assert warehouse.post(f"/api/documents/{doc['document_id']}/post", json={}).status_code == 200
        from api.db import connection, fetch_one
        with connection() as conn:
            row = fetch_one(conn, "SELECT posted_by_user_id FROM inventory_movement WHERE document_id=%s", (doc["document_id"],))
        self.assertIsNotNone(row["posted_by_user_id"], "API 过账流水应记录操作人用户外键")
        self.assertEqual(row["posted_by_user_id"], warehouse.user["user_id"])

    def test_orphan_party_rejected_by_fk(self) -> None:
        """新写入的 party 引用立即受外键约束（NOT VALID 不影响新行校验）。"""
        import psycopg2
        from tests.support.api_client import api_for
        admin_id = api_for("admin").user["user_id"]
        from api.db import connection
        with self.assertRaises(psycopg2.Error):
            with connection() as conn:
                from api.db import audit
                audit(conn, None, "TEST", "business_document", after={"doc_no": "FK-TEST"})
                with conn.cursor() as cur:
                    cur.execute("""INSERT INTO business_document(doc_type,doc_no,status,doc_date,party_type,party_id,created_by)
                                   VALUES ('PURCHASE_ORDER','FK-TEST-1','DRAFT',current_date,'CUSTOMER',99999999,%s)""", (admin_id,))


if __name__ == "__main__":
    unittest.main()
