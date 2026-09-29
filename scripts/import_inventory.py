#!/usr/bin/env python3
"""Read a mapped inventory workbook into safe, auditable dry-run records.

This command intentionally stops before posting ledger rows.  It produces JSON
and JSONL staging artifacts that can be reviewed and loaded into the SQL staging
tables.  The password is read from an environment variable and is never
written to output or logs; encrypted workbooks are decrypted into a temporary
directory which is removed on exit.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import struct
import tempfile
import unicodedata
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator


try:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    from cryptography.hazmat.primitives.hashes import SHA1
    from cryptography.hazmat.backends import default_backend
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - clear runtime diagnostic
    raise SystemExit("Requires bundled cryptography and openpyxl: %s" % exc)


SUPPORTED_MOVEMENT_TYPES = {
    "OPENING", "RECEIPT", "ISSUE_OTHER", "ISSUE_SALE", "ISSUE_CONSUMPTION",
    "ISSUE_GIFT", "ISSUE_SCRAP", "TRANSFER", "RETURN", "ADJUSTMENT",
}
ROLE_FIELDS = {
    "products": {"identifier", "name", "uom", "opening_quantity", "existing_quantity"},
    "movement": {"identifier", "name", "date", "quantity", "uom", "note", "destination", "serial"},
    "assets": {"name", "serial", "component_serial", "notes"},
    "stage_only": None,
}
TRANSFER_KEYWORDS = (
    ("ISSUE_SCRAP", ("报废", "作废")),
    ("ISSUE_GIFT", ("赠送", "送给", "赠")),
    ("ISSUE_CONSUMPTION", ("消耗", "使用", "领用")),
    ("ISSUE_SALE", ("售出", "销售", "出售", "卖出")),
    ("TRANSFER", ("备用", "借用", "暂借", "摆放", "调货", "调拨")),
)

_PHONE_RE = re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)|(?<!\d)0\d{2,3}[- ]?\d{7,8}(?!\d)")
_SECRET_KEY_RE = re.compile(r"(?:密码|口令|password|passwd|账号|用户名|user(name)?|login|token|secret)", re.I)
_PHONE_KEY_RE = re.compile(r"(?:电话|手机|手机号|联系电话|联系方式|客户电话|phone|mobile)", re.I)
_EXCLUDED_KEY_RE = re.compile(r"(?:密码|口令|password|passwd|账号|用户名|login|token|secret|电话|手机|手机号|联系电话|联系方式|客户电话|客户|付款|是否付款|金额|价款|成本|费用|价格|售价|销售价格|提成|commission|amount|price|cost|revenue|payment)", re.I)
_FINANCIAL_TEXT_RE = re.compile(r"(?:付款|已付|已收|发票|金额|价款|成本|费用|价格|售价|提成|佣金|commission|(?:amount|price|cost|revenue|payment)|[0-9０-９零一二三四五六七八九十百千万亿点]+\s*(?:千|万|元))", re.I)
_QTY_RE = re.compile(r"^[\s]*([+-]?(?:\d+(?:\.\d*)?|\.\d+))(?:\s*(?:个|件|台|只|套|盒|箱|包|支|米|ml|ML|pcs?|EA|盒装))?[\s]*$")
_SERIAL_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])([A-Za-z0-9]+(?:-[A-Za-z0-9]+)+|[A-Za-z]*\d[A-Za-z0-9]{5,})(?![A-Za-z0-9])")


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def normalize_identifier(value: Any) -> str | None:
    """Normalize an identifier without converting it to a number."""
    text = normalize_text(value)
    return text.casefold() if text else None


def parse_quantity(value: Any) -> tuple[float | None, str | None]:
    if value is None or normalize_text(value) is None:
        return None, "missing_quantity"
    if isinstance(value, bool):
        return None, "invalid_quantity"
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        match = _QTY_RE.match(normalize_text(value) or "")
        if not match:
            return None, "invalid_quantity"
        number = float(match.group(1))
    if number == 0:
        return number, "zero_quantity"
    if number < 0:
        return number, "negative_quantity"
    return number, None


def parse_date(value: Any) -> tuple[str | None, str | None]:
    if value is None or normalize_text(value) is None:
        return None, "missing_date"
    parsed: date | None = None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = normalize_text(value) or ""
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
    if parsed is None:
        return None, "invalid_date"
    if not (date(1900, 1, 1) <= parsed <= date(2200, 1, 1)):
        return parsed.isoformat(), "out_of_range_date"
    return parsed.isoformat(), None


def classify_transfer(text: Any, destination: Any = None) -> tuple[str, str]:
    haystack = " ".join(x for x in (normalize_text(text), normalize_text(destination)) if x).casefold()
    for movement_type, words in TRANSFER_KEYWORDS:
        if any(word.casefold() in haystack for word in words):
            if movement_type == "TRANSFER":
                return movement_type, "keyword_transfer"
            return movement_type, "keyword_terminal_outflow"
    return "REVIEW", "no_safe_transfer_semantics"


def redact_value(value: Any, key: str | None = None) -> Any:
    """Remove credentials and phone numbers from all emitted artifacts."""
    if key and (_SECRET_KEY_RE.search(key) or _PHONE_KEY_RE.search(key)):
        return "[REDACTED]"
    if isinstance(value, dict):
        return {str(k): redact_value(v, str(k)) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [redact_value(v, key) for v in value]
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric_text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
        if _PHONE_RE.fullmatch(numeric_text):
            return "[REDACTED_PHONE]"
        return value
    if isinstance(value, str):
        if _FINANCIAL_TEXT_RE.search(value):
            return "[REDACTED_FINANCIAL_OR_PERSONAL_NOTE]"
        return _PHONE_RE.sub("[REDACTED_PHONE]", value)
    return value


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def _read_u16(buf: bytes, off: int) -> int:
    return struct.unpack_from("<H", buf, off)[0]


def _read_u32(buf: bytes, off: int) -> int:
    return struct.unpack_from("<I", buf, off)[0]


class CompoundFile:
    """Small read-only OLE compound-file reader for the Office Standard format."""

    def __init__(self, data: bytes):
        if data[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
            raise ValueError("not an OLE compound file")
        self.data = data
        self.sector_size = 1 << _read_u16(data, 30)
        self.mini_sector_size = 1 << _read_u16(data, 32)
        self.first_dir = struct.unpack_from("<i", data, 48)[0]
        self.mini_cutoff = _read_u32(data, 56)
        self.first_mini_fat = struct.unpack_from("<i", data, 60)[0]
        self.mini_fat_count = _read_u32(data, 64)
        self.first_difat = struct.unpack_from("<i", data, 68)[0]
        self.difat_count = _read_u32(data, 72)
        difat = list(struct.unpack_from("<109i", data, 76))
        sid = self.first_difat
        for _ in range(self.difat_count):
            block = self.sector(sid)
            count = self.sector_size // 4 - 1
            difat.extend(struct.unpack_from("<%di" % count, block, 0)[:count])
            sid = struct.unpack_from("<i", block, count * 4)[0]
        self.fat: list[int] = []
        for sid in difat:
            if sid >= 0:
                block = self.sector(sid)
                self.fat.extend(struct.unpack("<%di" % (self.sector_size // 4), block))
        self.entries = self._directory()
        root = next((e for e in self.entries if e["type"] == 5), None)
        if not root:
            raise ValueError("OLE root storage missing")
        self.root_mini_stream = self._read_regular(root["start"], root["size"])
        self.mini_fat: list[int] = []
        if self.first_mini_fat >= 0:
            for sid in self._chain(self.first_mini_fat, self.mini_fat_count):
                self.mini_fat.extend(struct.unpack("<%di" % (self.sector_size // 4), self.sector(sid)))

    def sector(self, sid: int) -> bytes:
        start = (sid + 1) * self.sector_size
        return self.data[start : start + self.sector_size]

    def _chain(self, start: int, limit: int | None = None) -> Iterator[int]:
        seen: set[int] = set()
        sid = start
        count = 0
        while sid >= 0 and sid not in seen and sid not in (0xFFFFFFFE, 0xFFFFFFFF):
            yield sid
            seen.add(sid)
            count += 1
            if limit is not None and count >= limit:
                break
            if sid >= len(self.fat):
                break
            sid = self.fat[sid]

    def _read_regular(self, start: int, size: int) -> bytes:
        if start < 0 or size <= 0:
            return b""
        out = b"".join(self.sector(sid) for sid in self._chain(start))
        return out[:size]

    def _directory(self) -> list[dict[str, Any]]:
        raw = self._read_regular(self.first_dir, 0x7FFFFFFF)
        entries: list[dict[str, Any]] = []
        for off in range(0, len(raw) - 127, 128):
            name_len = _read_u16(raw, off + 64)
            name = raw[off : off + max(0, name_len - 2)].decode("utf-16le", "ignore") if name_len else ""
            entries.append({"name": name, "type": raw[off + 66], "start": struct.unpack_from("<i", raw, off + 116)[0], "size": struct.unpack_from("<Q", raw, off + 120)[0]})
        return entries

    def stream(self, name: str) -> bytes:
        entry = next((e for e in self.entries if e["name"] == name and e["type"] == 2), None)
        if not entry:
            raise KeyError(name)
        if entry["size"] < self.mini_cutoff:
            out = bytearray()
            sid = entry["start"]
            seen: set[int] = set()
            while sid >= 0 and sid not in seen and sid not in (0xFFFFFFFE, 0xFFFFFFFF):
                seen.add(sid)
                begin = sid * self.mini_sector_size
                out.extend(self.root_mini_stream[begin : begin + self.mini_sector_size])
                sid = self.mini_fat[sid]
            return bytes(out[: entry["size"]])
        return self._read_regular(entry["start"], entry["size"])


def _standard_key(password: str, salt: bytes, block: int, key_size: int) -> bytes:
    digest = SHA1()
    from cryptography.hazmat.primitives.hashes import Hash
    h = Hash(digest, backend=default_backend())
    h.update(salt + password.encode("utf-16le"))
    current = h.finalize()
    for i in range(50000):
        h = Hash(SHA1(), backend=default_backend())
        h.update(struct.pack("<I", i) + current)
        current = h.finalize()
    h = Hash(SHA1(), backend=default_backend())
    h.update(current + struct.pack("<I", block))
    hfinal = h.finalize()
    # ECMA-376 Standard derives the AES key through SHA1(H xor ipad/opad),
    # not by truncating Hfinal directly.
    ipad = bytes((b ^ 0x36) for b in hfinal) + bytes([0x36]) * (64 - len(hfinal))
    opad = bytes((b ^ 0x5C) for b in hfinal) + bytes([0x5C]) * (64 - len(hfinal))
    h = Hash(SHA1(), backend=default_backend()); h.update(ipad); x1 = h.finalize()
    h = Hash(SHA1(), backend=default_backend()); h.update(opad); x2 = h.finalize()
    return (x1 + x2)[:key_size]


def decrypt_standard_xlsx(source: Path, password: str, target: Path) -> None:
    ole = CompoundFile(source.read_bytes())
    info = ole.stream("EncryptionInfo")
    if _read_u16(info, 2) != 2:
        raise ValueError("only Office Standard encryption is supported")
    header_size = _read_u32(info, 8)
    header = info[12 : 12 + header_size]
    key_size_bits = _read_u32(header, 16)
    key_size = max(16, key_size_bits // 8)
    salt = info[12 + header_size + 4 : 12 + header_size + 20]
    encrypted_verifier = info[12 + header_size + 20 : 12 + header_size + 36]
    # Standard verifier layout: saltSize(4), salt(16), verifier(16),
    # verifierHashSize(4), encryptedVerifierHash(32).
    encrypted_hash = info[12 + header_size + 40 : 12 + header_size + 72]
    aes = lambda key: Cipher(algorithms.AES(key), modes.ECB(), backend=default_backend()).decryptor()
    verifier_key = _standard_key(password, salt, 0, key_size)
    verifier = aes(verifier_key).update(encrypted_verifier)
    verifier_hash = aes(verifier_key).update(encrypted_hash)
    from cryptography.hazmat.primitives.hashes import Hash
    h = Hash(SHA1(), backend=default_backend()); h.update(verifier); expected = h.finalize()
    if expected != verifier_hash[:20]:
        raise ValueError("incorrect workbook password")
    package = ole.stream("EncryptedPackage")
    plain_size = struct.unpack_from("<I", package, 0)[0]
    encrypted = package[8:]
    out = bytearray()
    # Standard CryptoAPI encrypts the complete package with one AES-ECB key;
    # the 4,096-byte chunking used by Agile encryption does not apply here.
    decryptor = aes(verifier_key)
    out.extend(decryptor.update(encrypted) + decryptor.finalize())
    target.write_bytes(bytes(out[:plain_size]))


@contextmanager
def decrypted_workbook(path: Path, password: str) -> Iterator[Path]:
    if path.suffix.lower() != ".xlsx" or path.read_bytes()[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
        yield path
        return
    temp_dir = Path(tempfile.mkdtemp(prefix="mini-erp-import-"))
    decrypted = temp_dir / "decrypted.xlsx"
    try:
        decrypt_standard_xlsx(path, password, decrypted)
        yield decrypted
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def _extract_serial_tokens(values: Iterable[Any]) -> list[str]:
    found: set[str] = set()
    for value in values:
        for match in _SERIAL_TOKEN_RE.finditer(normalize_text(value) or ""):
            token = match.group(1)
            if not re.fullmatch(r"1[3-9]\d{9}", token):
                found.add(token)
    return sorted(found)


def _mapping_error(message: str) -> ValueError:
    return ValueError(f"导入映射无效：{message}")


def load_mapping(path: Path) -> dict[str, Any]:
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise _mapping_error(f"无法读取 JSON 文件（{exc}）") from exc
    if not isinstance(config, dict) or isinstance(config.get("version"), bool) or config.get("version") != 1:
        raise _mapping_error("version 必须为 1")
    entries = config.get("sheets")
    if not isinstance(entries, list) or not entries:
        raise _mapping_error("sheets 必须是非空数组")

    seen_names: set[str] = set()
    product_sheets = 0
    normalized_entries: list[dict[str, Any]] = []
    for index, entry in enumerate(entries):
        if not isinstance(entry, dict):
            raise _mapping_error(f"sheets[{index}] 必须是对象")
        name = normalize_text(entry.get("name"))
        role = entry.get("role")
        if not name:
            raise _mapping_error(f"sheets[{index}].name 不能为空")
        if name in seen_names:
            raise _mapping_error(f"工作表重复配置：{name}")
        seen_names.add(name)
        if role not in ROLE_FIELDS:
            raise _mapping_error(f"工作表 {name} 的 role 不受支持")
        if role == "products":
            product_sheets += 1

        header_row = entry.get("header_row", 1)
        if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
            raise _mapping_error(f"工作表 {name} 的 header_row 必须是正整数")
        columns = entry.get("columns")
        if not isinstance(columns, dict):
            raise _mapping_error(f"工作表 {name} 的 columns 必须是对象")
        allowed_fields = ROLE_FIELDS[role]
        if allowed_fields is not None:
            unknown = set(columns) - allowed_fields
            if unknown:
                raise _mapping_error(f"工作表 {name} 含不支持的字段：{', '.join(sorted(map(str, unknown)))}")
        clean_columns: dict[str, str] = {}
        for field, source_header in columns.items():
            if not isinstance(field, str) or not field.strip():
                raise _mapping_error(f"工作表 {name} 含空标准字段名")
            if not isinstance(source_header, str):
                raise _mapping_error(f"工作表 {name} 的源列标题必须是字符串")
            header = normalize_text(source_header)
            if not header:
                raise _mapping_error(f"工作表 {name} 的源列标题不能为空")
            if role != "stage_only" and (_EXCLUDED_KEY_RE.search(header) or _PHONE_KEY_RE.search(header)):
                raise _mapping_error(f"工作表 {name} 不能把敏感列映射到业务字段：{header}")
            clean_columns[field] = header
        if role == "products" and "name" not in clean_columns:
            raise _mapping_error(f"产品工作表 {name} 必须映射 name")
        if role == "movement":
            movement_type = entry.get("movement_type")
            if movement_type not in SUPPORTED_MOVEMENT_TYPES:
                raise _mapping_error(f"工作表 {name} 的 movement_type 不受支持")
            if not ({"identifier", "name"} & set(clean_columns)):
                raise _mapping_error(f"流水工作表 {name} 至少要映射 identifier 或 name")
        else:
            movement_type = None
            if "movement_type" in entry:
                raise _mapping_error(f"只有 movement 工作表可设置 movement_type：{name}")
        normalized_entries.append({
            "name": name,
            "role": role,
            "header_row": header_row,
            "movement_type": movement_type,
            "columns": clean_columns,
        })
    if product_sheets != 1:
        raise _mapping_error("必须且只能配置一个 products 工作表")
    return {"version": 1, "sheets": normalized_entries}


def _mapped_rows(wb: Any, entry: dict[str, Any]) -> Iterator[tuple[int, dict[str, Any], dict[str, Any]]]:
    sheet_name = entry["name"]
    if sheet_name not in wb.sheetnames:
        raise _mapping_error(f"工作簿缺少工作表：{sheet_name}")
    ws = wb[sheet_name]
    header_row = entry["header_row"]
    if header_row > ws.max_row:
        raise _mapping_error(f"工作表 {sheet_name} 不存在第 {header_row} 行表头")
    headers = [normalize_text(value) for value in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    header_indices: dict[str, int] = {}
    for index, header in enumerate(headers):
        if not header:
            continue
        if header in header_indices:
            raise _mapping_error(f"工作表 {sheet_name} 的表头重复：{header}")
        header_indices[header] = index
    field_indices: dict[str, tuple[int, str]] = {}
    for field, source_header in entry["columns"].items():
        if source_header not in header_indices:
            raise _mapping_error(f"工作表 {sheet_name} 缺少映射列：{source_header}")
        field_indices[field] = (header_indices[source_header], source_header)

    for row_number, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if not any(value is not None for value in row):
            continue
        raw: dict[str, Any] = {}
        safe: dict[str, Any] = {}
        for field, (column_index, source_header) in field_indices.items():
            value = row[column_index] if column_index < len(row) else None
            raw[field] = value
            if not (_EXCLUDED_KEY_RE.search(source_header) or _PHONE_KEY_RE.search(source_header) or _EXCLUDED_KEY_RE.search(field) or _PHONE_KEY_RE.search(field)):
                safe[field] = redact_value(json_value(value), f"{source_header} {field}")
        yield row_number, raw, safe


def _extract_mapped(wb: Any, mapping: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], Counter]:
    observations: list[dict[str, Any]] = []
    movements: list[dict[str, Any]] = []
    assets: list[dict[str, Any]] = []
    staged: list[dict[str, Any]] = []
    counts: Counter = Counter()

    for entry in mapping["sheets"]:
        role = entry["role"]
        for row_number, raw, safe in _mapped_rows(wb, entry):
            if role == "products":
                identifier = safe.get("identifier")
                name = safe.get("name")
                if normalize_text(identifier) is None and normalize_text(name) is None:
                    continue
                observations.append({
                    "sheet": entry["name"],
                    "source_row_number": row_number,
                    "identifier_raw": json_value(identifier),
                    "identifier_normalized": normalize_identifier(identifier),
                    "name_raw": json_value(name),
                    "name_normalized": normalize_identifier(name),
                    "source": safe,
                })
                counts["products"] += 1
            elif role == "movement":
                identifier = safe.get("identifier")
                name = safe.get("name")
                date_raw = safe.get("date")
                quantity_raw = safe.get("quantity")
                if all(normalize_text(value) is None for value in (identifier, name, date_raw, quantity_raw)):
                    continue
                quantity, quantity_issue = parse_quantity(raw.get("quantity"))
                movement_date, date_issue = parse_date(raw.get("date"))
                movement_type = entry["movement_type"]
                if movement_type == "TRANSFER":
                    movement_type, method = classify_transfer(safe.get("note"), safe.get("destination"))
                else:
                    method = "mapping_movement_type"
                issues = [issue for issue in (quantity_issue, date_issue) if issue]
                movements.append({
                    "sheet": entry["name"],
                    "source_row_number": row_number,
                    "source": safe,
                    "identifier_raw": json_value(identifier),
                    "identifier_normalized": normalize_identifier(identifier),
                    "name_raw": json_value(name),
                    "name_normalized": normalize_identifier(name),
                    "movement_type_candidate": movement_type,
                    "classification_method": method,
                    "movement_date_raw": json_value(date_raw),
                    "movement_date": movement_date,
                    "quantity_raw": json_value(safe.get("quantity")),
                    "quantity": quantity,
                    "serial_candidates": _extract_serial_tokens([safe.get("serial")]),
                    "data_quality_issues": issues,
                })
                counts["movements"] += 1
            elif role == "assets":
                if not any(normalize_text(value) for value in safe.values()):
                    continue
                serial_candidates: set[str] = set()
                for field in ("serial", "component_serial"):
                    value = safe.get(field)
                    normalized = normalize_text(value)
                    if not normalized:
                        continue
                    tokens = _extract_serial_tokens([normalized])
                    serial_candidates.update(tokens or [normalized])
                assets.append({
                    "sheet": entry["name"],
                    "source_row_number": row_number,
                    "source": safe,
                    "serial_candidates": sorted(serial_candidates),
                    "data_quality_issues": [],
                })
                counts["assets"] += 1
            else:
                if not safe:
                    continue
                staged.append({
                    "sheet": entry["name"],
                    "source_row_number": row_number,
                    "source": safe,
                    "data_quality_issues": ["stage_only"],
                })
                counts["stage_only"] += 1

    return observations, movements, assets, staged, counts


def build_report(input_path: Path, password: str, mapping_path: Path) -> dict[str, Any]:
    mapping = load_mapping(mapping_path)
    source_bytes = input_path.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    with decrypted_workbook(input_path, password) as workbook_path:
        wb = load_workbook(workbook_path, data_only=True, read_only=False)
        mapped_names = {entry["name"] for entry in mapping["sheets"]}
        for entry in mapping["sheets"]:
            if entry["name"] not in wb.sheetnames:
                raise _mapping_error(f"工作簿缺少工作表：{entry['name']}")
        unmapped_sheets = []
        for sheet_name in wb.sheetnames:
            if sheet_name not in mapped_names:
                ws = wb[sheet_name]
                non_empty_rows = sum(1 for row in ws.iter_rows(values_only=True) if any(value is not None for value in row))
                unmapped_sheets.append({"name": sheet_name, "row_count": non_empty_rows})
        observations, movements, assets, staged, counts = _extract_mapped(wb, mapping)
        names_by_identifier: defaultdict[str, set[str]] = defaultdict(set)
        for row in observations:
            if row.get("identifier_normalized"):
                names_by_identifier[row["identifier_normalized"]].add(row.get("name_normalized") or "")
        resolution_counts: Counter = Counter()
        for row in observations:
            identifier = row.get("identifier_normalized")
            if not identifier:
                row["resolution_status"] = "pending_review"
                row["match_method"] = "missing_identifier"
            elif len(names_by_identifier[identifier]) == 1:
                row["resolution_status"] = "candidate_exact"
                row["match_method"] = "identifier_name_exact"
            else:
                row["resolution_status"] = "pending_review"
                row["match_method"] = "identifier_collision"
            resolution_counts[row["resolution_status"]] += 1
        asset_serials = {serial for row in assets for serial in row.get("serial_candidates", [])}
        movement_serials = {serial for row in movements for serial in row.get("serial_candidates", [])}
        summary = {
            "source_file_name": input_path.name,
            "source_sha256": source_hash,
            "encrypted_source": source_bytes[:8] == bytes.fromhex("D0CF11E0A1B11AE1"),
            "sheet_names": list(wb.sheetnames),
            "unmapped_sheets": unmapped_sheets,
            "counts": dict(counts),
            "product_resolution": {
                "candidate_exact": resolution_counts.get("candidate_exact", 0),
                "pending_review": resolution_counts.get("pending_review", 0),
                "identifier_groups": len(names_by_identifier),
                "collision_groups": sum(1 for names in names_by_identifier.values() if len(names) > 1),
            },
            "serial_linkage": {
                "asset_serial_candidates": len(asset_serials),
                "movement_serial_candidates": len(movement_serials),
                "exact_candidate_intersection": len(asset_serials & movement_serials),
            },
            "posting_policy": "dry-run only; review/approval required before inventory_movement status=posted",
            "product_resolution_policy": "internal product_id; source identifier is a non-unique candidate",
            "movement_policy": "ambiguous transfers remain REVIEW",
            "security_policy": "only mapped columns are exported; sensitive fields are excluded or redacted",
        }
        wb.close()
        return {
            "summary": summary,
            "product_observations": observations,
            "movement_candidates": movements,
            "asset_observations": assets,
            "stage_only": staged,
        }


def write_report(report: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "report.json").write_text(json.dumps(redact_value(report), ensure_ascii=False, indent=2), encoding="utf-8")
    for key in ("product_observations", "movement_candidates", "asset_observations", "stage_only"):
        with (output_dir / f"{key}.jsonl").open("w", encoding="utf-8") as handle:
            for row in report[key]:
                handle.write(json.dumps(redact_value(row), ensure_ascii=False) + "\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--mapping", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--password-env", default="IMPORT_WORKBOOK_PASSWORD", help="environment variable containing workbook password")
    args = parser.parse_args(argv)
    password = os.environ.get(args.password_env)
    if not password:
        parser.error(f"environment variable {args.password_env} is empty")
    if not args.input.exists():
        parser.error(f"找不到工作簿：{args.input}")
    if not args.mapping.is_file():
        parser.error(f"找不到导入映射：{args.mapping}")
    report = build_report(args.input, password, args.mapping)
    write_report(report, args.output)
    print(json.dumps(report["summary"], ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
