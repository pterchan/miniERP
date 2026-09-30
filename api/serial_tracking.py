"""货品唯一编码（SN/UUID）流向追踪：激活并复用 001 迁移的休眠 asset 资产域。

每个启用 SN 的货品，其单件以 asset 行 + asset_identifier(product_serial) 建档，
流向通过 asset_event 记录，并与 inventory_movement 经 inventory_movement_asset 关联。
所有写库均在同一事务内先 audit() 再 INSERT——asset 系列表的 require_audit_context
触发器在 002 已挂载，事务内每条写入前必须存在对应 target_table 的 audit_event。

登记可选（软约束）：serial_numbers 为空则整体跳过；一旦填写即按数据完整性硬校验。
"""

from __future__ import annotations

import csv
import io
import re
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, Response, UploadFile
from openpyxl import load_workbook

from .db import audit, connection, fetch_all, fetch_one
from .export import MAX_EXPORT_ROWS, export_response
from .helpers import _condition_id, _normalize_identifier, _status_id
from .list_params import clamp_page, clamp_page_size, like_escape, parse_filters, parse_sort
from .permissions import _csrf, require_user
from .schemas import SerialParseIn

SN_IDENT_TYPE = "product_serial"
MAX_IMPORT_BYTES = 5 * 1024 * 1024  # 导入文件 ≤5MB

router = APIRouter(prefix="/api/serial-ledger", tags=["serial"])


# ---------------------------------------------------------------------------
# SN 解析与归一化
# ---------------------------------------------------------------------------

def normalize_sn(value: str) -> str:
    """归一化序列号：NFKC + 折叠空白 + casefold，与货品编号一致。"""
    return _normalize_identifier(value)


def parse_serial_block(raw: str | None) -> list[str]:
    """按 换行/逗号/分号/制表符 切分 → 归一化 → 去空 → 去重保序。"""
    parts = re.split(r"[\s,;，；、]+", raw or "")
    seen: list[str] = []
    for p in parts:
        n = normalize_sn(p)
        if n and n not in seen:
            seen.append(n)
    return seen


def _extract_serials_from_file(data: bytes, filename: str) -> list[str]:
    """从 xlsx / csv 第一列提取 SN 原始值（纯解析，不写库）。"""
    name = (filename or "").lower()
    raw_values: list[str] = []
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        # openpyxl 不支持旧版 BIFF（.xls）；明确提示而不是解析异常 500
        raise HTTPException(status_code=422, detail="不支持旧版 .xls，请用 Excel 另存为 .xlsx 或导出 .csv 后重试")
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True)
        except Exception as exc:
            raise HTTPException(status_code=422, detail="无法解析的表格文件，请确认是有效的 .xlsx") from exc
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            if row and row[0] is not None and str(row[0]).strip():
                raw_values.append(str(row[0]))
    else:  # csv / txt：每行一个
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            if row and row[0].strip():
                raw_values.append(row[0].strip())
    return parse_serial_block("\n".join(raw_values))


# ---------------------------------------------------------------------------
# 资产域写入助手（每个写库存事务内先 audit() 再 INSERT）
# ---------------------------------------------------------------------------

def _asset_by_sn(conn: Any, product_id: int, sn: str) -> dict[str, Any] | None:
    """按 (product, SN) 查已建档资产及其**当前状态**；未建档返回 None。

    当前状态必须取自 v_asset_current_state（最新事件推导）：asset.status_id 只是
    建档初值，出库/回库只追加 asset_event，不改 asset 行——直接读列会把已出库
    的 SN 永远当成在库（重复出库放行、退货回库被拒）。
    """
    return fetch_one(
        conn,
        """SELECT s.asset_id, s.status_id, s.condition_id, s.status_code,
                  s.current_location_id, s.latest_event_type, im.uom_id AS current_uom_id
             FROM v_asset_current_state s
             JOIN asset_identifier ai ON ai.asset_id = s.asset_id
             LEFT JOIN inventory_movement im ON im.inventory_movement_id=s.latest_inventory_movement_id
            WHERE ai.namespace=%s AND ai.identifier_type=%s AND ai.value_normalized=%s
              AND ai.is_verified AND ai.is_exclusive""",
        (f"{SN_IDENT_TYPE}.{product_id}", SN_IDENT_TYPE, normalize_sn(sn)),
    )


def _require_in_stock(asset: dict[str, Any] | None, sn: str, action: str, *,
                      location_id: int | None, condition_id: int, uom_id: int | None) -> int:
    """出库/调拨/清点出库共用的在库校验；返回 asset_id。"""
    if not asset:
        raise HTTPException(status_code=422, detail=f"SN {sn} 未在库中登记，无法{action}")
    if asset["status_code"] != "active":
        raise HTTPException(status_code=422, detail=f"SN {sn} 当前状态为 {asset['status_code']}（不在库），无法{action}")
    if asset["current_location_id"] != location_id:
        raise HTTPException(status_code=422,
                            detail=f"SN {sn} 当前库位为 {asset['current_location_id']}，与{action}来源库位 {location_id} 不一致")
    if asset["condition_id"] != condition_id:
        raise HTTPException(status_code=422, detail=f"SN {sn} 当前成色与{action}明细成色不一致")
    if uom_id is None or asset.get("current_uom_id") is None:
        raise HTTPException(status_code=422, detail=f"SN {sn} 缺少可核对的当前库存单位，无法{action}；请先核对原入库记录")
    if asset["current_uom_id"] != uom_id:
        raise HTTPException(status_code=422, detail=f"SN {sn} 当前库存单位与{action}明细单位不一致；系统不进行单位换算")
    return int(asset["asset_id"])


def _create_asset(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
                  product_id: int, sn: str, condition_id: int) -> int:
    """为 (product, SN) 建档 asset + product_serial 标识，返回 asset_id。"""
    product = fetch_one(conn, "SELECT display_name, manufacturer, specification FROM product WHERE product_id=%s", (product_id,))
    audit(conn, user, "CREATE", "asset",
          after={"product_id": product_id, "sn": sn, "asset_type": "serialized_product"},
          request_id=req_meta["request_id"], ip_address=req_meta["ip_address"], user_agent=req_meta["user_agent"])
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO asset(product_id,asset_type,manufacturer,model,status_id,condition_id,notes)
                     VALUES (%s,'serialized_product',%s,%s,%s,%s,%s) RETURNING asset_id""",
                    (product_id, product["manufacturer"] if product else None,
                     product["specification"] if product else None,
                     _status_id(conn, "active"), condition_id, f"SN 追踪自动建档 {sn}"))
        asset_id = cur.fetchone()[0]
    audit(conn, user, "CREATE", "asset_identifier",
          after={"asset_id": asset_id, "product_id": product_id, "sn": sn, "identifier_type": SN_IDENT_TYPE},
          request_id=req_meta["request_id"], ip_address=req_meta["ip_address"], user_agent=req_meta["user_agent"])
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO asset_identifier(asset_id,identifier_type,namespace,value_raw,value_normalized,is_primary,is_verified,is_exclusive)
                     VALUES (%s,%s,%s,%s,%s,TRUE,TRUE,TRUE)""",
                    (asset_id, SN_IDENT_TYPE, f"{SN_IDENT_TYPE}.{product_id}", sn, normalize_sn(sn)))
    return asset_id


def _write_event(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
                 asset_id: int, movement_id: int, event_type: str, status_code: str,
                 condition_id: int, from_loc: int | None, to_loc: int | None, notes: str | None = None) -> None:
    audit(conn, user, "CREATE", "asset_event",
          after={"asset_id": asset_id, "event_type": event_type, "movement_id": movement_id},
          request_id=req_meta["request_id"], ip_address=req_meta["ip_address"], user_agent=req_meta["user_agent"])
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO asset_event(asset_id,event_type,event_date,status_id,condition_id,from_location_id,to_location_id,inventory_movement_id,notes)
                     VALUES (%s,%s,current_date,%s,%s,%s,%s,%s,%s)""",
                    (asset_id, event_type, _status_id(conn, status_code), condition_id,
                     from_loc, to_loc, movement_id, notes))


def _link(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
          movement_id: int, asset_id: int) -> None:
    audit(conn, user, "CREATE", "inventory_movement_asset",
          after={"movement_id": movement_id, "asset_id": asset_id},
          request_id=req_meta["request_id"], ip_address=req_meta["ip_address"], user_agent=req_meta["user_agent"])
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO inventory_movement_asset(inventory_movement_id,asset_id,asset_role,quantity)
                     VALUES (%s,%s,'primary',1)""", (movement_id, asset_id))


def _validate_integer_quantity(quantity: Decimal, count: int) -> None:
    if quantity != quantity.to_integral_value():
        raise HTTPException(status_code=422, detail="填了序列号时数量必须为整数（序列号按件计）")
    if int(quantity) != count:
        raise HTTPException(status_code=422, detail=f"序列号数量（{count}）与单据数量（{int(quantity)}）不一致")


# ---------------------------------------------------------------------------
# 三处写库存路径的 SN hook
# ---------------------------------------------------------------------------

def apply_line_serials(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
                       movement_id: int, product_id: int, serial_numbers: list[str] | None,
                       stock_effect: str, source: int | None, dest: int | None,
                       quantity: Decimal, condition_id: int | None, uom_id: int | None = None) -> None:
    """单据过账后按行登记 SN：IN 建档/received，OUT issued，TRANSFER transferred。"""
    sns = [normalize_sn(s) for s in (serial_numbers or []) if (s or "").strip()]
    if not sns:
        return  # 软约束：未填 SN 直接跳过
    product = fetch_one(conn, "SELECT display_name, serialized FROM product WHERE product_id=%s", (product_id,))
    if not product or not product["serialized"]:
        raise HTTPException(status_code=422, detail=f"货品 {product['display_name'] if product else product_id} 未启用序列号追踪，但填了 SN")
    if len(sns) != len(set(sns)):
        raise HTTPException(status_code=422, detail="同一行内存在重复 SN")
    _validate_integer_quantity(quantity, len(sns))
    cond_id = _condition_id(conn, condition_id)
    for sn in sns:
        asset = _asset_by_sn(conn, product_id, sn)
        if stock_effect == "IN":
            if asset and asset["status_code"] == "active":
                raise HTTPException(status_code=422, detail=f"SN {sn} 已登记且未出库，无法重复入库")
            asset_id = asset["asset_id"] if asset else _create_asset(conn, user, req_meta, product_id, sn, cond_id)
            _write_event(conn, user, req_meta, asset_id, movement_id, "received", "active", cond_id, None, dest)
        elif stock_effect == "OUT":
            asset_id = _require_in_stock(asset, sn, "出库", location_id=source, condition_id=cond_id, uom_id=uom_id)
            _write_event(conn, user, req_meta, asset_id, movement_id, "issued", "retired",
                         cond_id, source, dest)
        elif stock_effect == "TRANSFER":
            asset_id = _require_in_stock(asset, sn, "调拨", location_id=source, condition_id=cond_id, uom_id=uom_id)
            _write_event(conn, user, req_meta, asset_id, movement_id, "transferred", "active",
                         cond_id, source, dest)
        else:
            continue
        _link(conn, user, req_meta, movement_id, asset_id)


def _event_location(event: dict[str, Any]) -> int | None:
    """与 v_asset_current_state 使用相同的事件库位口径。"""
    if event["event_type"] in ("issued", "retired", "lost"):
        return event["to_location_id"]
    return event["to_location_id"] or event["from_location_id"]


def validate_reverse_movement_serials(conn: Any, orig_movement: dict[str, Any], *,
                                      ignore_reversal_movement_id: int | None = None) -> list[dict[str, Any]]:
    """在插入反向流水前校验 SN，并捕获反向事件应恢复的状态。

    后续原流水与其已过账红冲事件相互抵消；最新未抵消事件必须是待红冲流水。
    因此允许按逆序撤销，却拒绝仅仅绕回同一状态的后续出入库或调拨。
    """
    links = fetch_all(conn, "SELECT asset_id FROM inventory_movement_asset WHERE inventory_movement_id=%s",
                      (orig_movement["inventory_movement_id"],))
    s = orig_movement.get("source_location_id")
    d = orig_movement.get("destination_location_id")
    cond = _condition_id(conn, orig_movement.get("condition_id"))
    prepared: list[dict[str, Any]] = []
    for link in links:
        asset_id = link["asset_id"]
        asset = fetch_one(conn, """SELECT status_code,condition_id,current_location_id
                                    FROM v_asset_current_state WHERE asset_id=%s""", (asset_id,))
        events = fetch_all(conn, """SELECT ae.*,rs.code AS status_code,
                                          COALESCE(ae.condition_id,a.condition_id) AS effective_condition_id
                                     FROM asset_event ae JOIN asset a ON a.asset_id=ae.asset_id
                                     LEFT JOIN inventory_movement im ON im.inventory_movement_id=ae.inventory_movement_id
                                     LEFT JOIN record_status ims ON ims.status_id=im.status_id
                                     LEFT JOIN record_status rs ON rs.status_id=COALESCE(ae.status_id,a.status_id)
                                    WHERE ae.asset_id=%s
                                      AND (im.inventory_movement_id IS NULL OR
                                           (im.reversal_of_movement_id IS NULL AND ims.code='posted'))
                                      AND NOT EXISTS (
                                          SELECT 1 FROM inventory_movement rev
                                          JOIN record_status revs ON revs.status_id=rev.status_id AND revs.code='posted'
                                          WHERE rev.reversal_of_movement_id=ae.inventory_movement_id
                                            AND (%s IS NULL OR rev.inventory_movement_id<>%s))
                                    ORDER BY ae.event_date DESC,ae.asset_event_id DESC LIMIT 1""",
                           (asset_id, ignore_reversal_movement_id, ignore_reversal_movement_id))
        if not events or events[0]["inventory_movement_id"] != orig_movement["inventory_movement_id"]:
            raise HTTPException(status_code=422,
                                detail=f"红冲失败：SN 资产（id={asset_id}）存在未冲销的后续事件；请先按逆序红冲后续单据")
        current = events[0]
        if (not asset or asset["status_code"] != current["status_code"]
                or asset["condition_id"] != current["effective_condition_id"]
                or asset["current_location_id"] != _event_location(current)
                or asset["condition_id"] != cond):
            raise HTTPException(status_code=422,
                                detail=f"红冲失败：SN 资产（id={asset_id}）当前状态、库位或成色与原流水不一致；请先核对后续单据")
        # 恢复的是原事件发生前的真实状态，需包括当时已发生的红冲事件。
        # 有效事件仅用于撤销顺序检查，不能拿它的前一项替代真实历史前态。
        previous = fetch_one(conn, """SELECT ae.*,rs.code AS status_code,
                                            COALESCE(ae.condition_id,a.condition_id) AS effective_condition_id
                                       FROM asset_event ae JOIN asset a ON a.asset_id=ae.asset_id
                                       LEFT JOIN record_status rs ON rs.status_id=COALESCE(ae.status_id,a.status_id)
                                      WHERE ae.asset_id=%s AND (ae.event_date,ae.asset_event_id)<(%s,%s)
                                      ORDER BY ae.event_date DESC,ae.asset_event_id DESC LIMIT 1""",
                             (asset_id, current["event_date"], current["asset_event_id"]))
        was_in = s is None and d is not None
        restore_location = _event_location(previous) if previous else (None if was_in else s)
        from_location = d
        if from_location == restore_location:
            from_location = None
        prepared.append({
            "asset_id": asset_id,
            "event_type": "issued" if was_in else ("received" if d is None else "transferred"),
            "status_code": previous["status_code"] if previous else ("retired" if was_in else "active"),
            "condition_id": previous["effective_condition_id"] if previous else cond,
            "from_location_id": from_location,
            "to_location_id": restore_location,
        })
    return prepared


def reverse_movement_serials(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
                             orig_movement: dict[str, Any], rev_movement_id: int, *,
                             validated_assets: list[dict[str, Any]] | None = None) -> None:
    """写入已预校验的 SN 反向事件；兼容旧调用时忽略刚插入的反向流水。"""
    if validated_assets is None:
        validated_assets = validate_reverse_movement_serials(
            conn, orig_movement, ignore_reversal_movement_id=rev_movement_id)
    for asset in validated_assets:
        _write_event(conn, user, req_meta, asset["asset_id"], rev_movement_id,
                     asset["event_type"], asset["status_code"], asset["condition_id"],
                     asset["from_location_id"], asset["to_location_id"])
        _link(conn, user, req_meta, rev_movement_id, asset["asset_id"])


def apply_adjustment_serials(conn: Any, user: dict[str, Any] | None, req_meta: dict[str, Any],
                             movement_id: int, product: dict[str, Any],
                             serial_numbers: list[str] | None, delta: Decimal,
                             condition_id: int, location_id: int, uom_id: int | None = None) -> None:
    """快速清点（adjust_inventory）：delta>0 建档 received，delta<0 出库 issued。"""
    sns = [normalize_sn(s) for s in (serial_numbers or []) if (s or "").strip()]
    if not sns:
        return
    if not product["serialized"]:
        raise HTTPException(status_code=422, detail=f"货品 {product['display_name']} 未启用序列号追踪，但填了 SN")
    if len(sns) != len(set(sns)):
        raise HTTPException(status_code=422, detail="存在重复 SN")
    if delta != delta.to_integral_value() or abs(int(delta)) != len(sns):
        raise HTTPException(status_code=422, detail=f"序列号数量（{len(sns)}）与清点差值（{delta}）不一致")
    cond_id = _condition_id(conn, condition_id)
    for sn in sns:
        asset = _asset_by_sn(conn, product["product_id"], sn)
        if delta > 0:
            if asset and asset["status_code"] == "active":
                raise HTTPException(status_code=422, detail=f"SN {sn} 已登记且未出库，无法重复入库")
            asset_id = asset["asset_id"] if asset else _create_asset(conn, user, req_meta, product["product_id"], sn, cond_id)
            _write_event(conn, user, req_meta, asset_id, movement_id, "received", "active", cond_id, None, location_id)
        else:
            asset_id = _require_in_stock(asset, sn, "清点出库", location_id=location_id, condition_id=cond_id, uom_id=uom_id)
            _write_event(conn, user, req_meta, asset_id, movement_id, "issued", "retired", cond_id, location_id, None)
        _link(conn, user, req_meta, movement_id, asset_id)


# ---------------------------------------------------------------------------
# 序列台账 API
# ---------------------------------------------------------------------------

_SERIAL_SORT = {
    "serial_number": "cs.serial_number",
    "product_name": "cs.product_name",
    "current_location_name": "cs.current_location_name",
    "status_code": "cs.status_code",
    "latest_event_date": "cs.latest_event_date",
}

_SERIAL_FILTERS = {
    "product_id": ("cs.product_id", ("eq",)),
    "location_id": ("cs.current_location_id", ("eq",)),
    "status_code": ("cs.status_code", ("eq", "in")),
    "serial_number": ("cs.serial_number", ("contains", "eq")),
}

_SERIAL_COLUMNS = [
    ("序列号", lambda r: r["serial_number"]),
    ("货品", lambda r: r["product_name"] or ""),
    ("当前库位", lambda r: r["current_location_name"] or ""),
    ("状态", lambda r: r["status_code"] or ""),
    ("成色", lambda r: r["condition_code"] or ""),
    ("最近事件", lambda r: r["latest_event_type"] or ""),
    ("最近日期", lambda r: r["latest_event_date"] or ""),
]


def _serial_where(q: str, f: list[str]) -> tuple[str, list[Any]]:
    clauses: list[str] = []
    params: list[Any] = []
    if q.strip():
        clauses.append("cs.serial_number ILIKE %s")
        params.append("%" + like_escape(normalize_sn(q)) + "%")
    try:
        fw, fp = parse_filters(f, _SERIAL_FILTERS)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from None
    clauses.extend(fw)
    params.extend(fp)
    return (" AND ".join(clauses) if clauses else "TRUE"), params


@router.get("")
def serial_ledger(q: str = "", page: int = 1, page_size: int = 30, sort: str = "", order: str = "asc",
                  f: list[str] = Query(default=[]), user: dict[str, Any] = Depends(require_user)) -> dict[str, Any]:
    page = clamp_page(page)
    page_size = clamp_page_size(page_size, cap=500)
    where, params = _serial_where(q, f)
    order_by = parse_sort(sort, order, _SERIAL_SORT, "serial_number")
    with connection() as conn:
        rows = fetch_all(conn, f"SELECT * FROM v_serial_ledger cs WHERE {where} ORDER BY {order_by} LIMIT %s OFFSET %s",
                         tuple(params + [page_size, (page - 1) * page_size]))
        total = fetch_one(conn, f"SELECT count(*) AS n FROM v_serial_ledger cs WHERE {where}", tuple(params))
    return {"items": rows, "page": page, "page_size": page_size, "total": int(total["n"])}


@router.get("/export")
def serial_ledger_export(q: str = "", sort: str = "", order: str = "asc", f: list[str] = Query(default=[]),
                         fmt: str = "xlsx", user: dict[str, Any] = Depends(require_user)) -> Response:
    where, params = _serial_where(q, f)
    order_by = parse_sort(sort, order, _SERIAL_SORT, "serial_number")
    with connection() as conn:
        rows = fetch_all(conn, f"SELECT * FROM v_serial_ledger cs WHERE {where} ORDER BY {order_by} LIMIT %s", tuple(params + [MAX_EXPORT_ROWS + 1]))
    return export_response(rows, _SERIAL_COLUMNS, "序列台账", fmt)


@router.get("/{asset_id}")
def serial_asset(asset_id: int, user: dict[str, Any] = Depends(require_user)) -> dict[str, Any]:
    with connection() as conn:
        row = fetch_one(conn, "SELECT * FROM v_asset_current_state WHERE asset_id=%s", (asset_id,))
        if not row:
            raise HTTPException(status_code=404, detail="序列号不存在")
        row["events"] = fetch_all(conn, """SELECT ae.asset_event_id, ae.event_type, ae.event_date, ae.notes,
                                                  rs.code AS status_code, ic.code AS condition_code,
                                                  sl.name AS from_location_name, dl.name AS to_location_name,
                                                  im.inventory_movement_id, mt.code AS movement_type_code,
                                                  d.doc_no, d.doc_type, d.status AS doc_status
                                             FROM asset_event ae
                                             LEFT JOIN record_status rs ON rs.status_id = ae.status_id
                                             LEFT JOIN inventory_condition ic ON ic.condition_id = ae.condition_id
                                             LEFT JOIN location sl ON sl.location_id = ae.from_location_id
                                             LEFT JOIN location dl ON dl.location_id = ae.to_location_id
                                             LEFT JOIN inventory_movement im ON im.inventory_movement_id = ae.inventory_movement_id
                                             LEFT JOIN movement_type mt ON mt.movement_type_id = im.movement_type_id
                                             LEFT JOIN business_document d ON d.document_id = im.document_id
                                            WHERE ae.asset_id=%s
                                            ORDER BY ae.event_date DESC, ae.asset_event_id DESC""", (asset_id,))
        row["identifiers"] = fetch_all(conn, "SELECT identifier_type,value_raw FROM asset_identifier WHERE asset_id=%s ORDER BY asset_identifier_id", (asset_id,))
    return row


@router.post("/parse")
def parse_serials(payload: SerialParseIn, request: Request = None,
                  user: dict[str, Any] = Depends(require_user)) -> dict[str, Any]:
    """解析一段 SN 文本并标注是否已在库中登记，供出库前预检。纯读取不写库。"""
    _csrf(request)
    sns = parse_serial_block(payload.text)
    with connection() as conn:
        items: list[dict[str, Any]] = []
        if sns:
            # 一次批量查询代替逐 SN 往返（文本最长 2 万字符可拆数千 SN）
            rows = fetch_all(conn, """SELECT ai.value_normalized AS sn, s.status_code
                                        FROM asset_identifier ai
                                        JOIN v_asset_current_state s ON s.asset_id = ai.asset_id
                                       WHERE ai.namespace=%s AND ai.identifier_type=%s
                                         AND ai.value_normalized = ANY(%s)
                                         AND ai.is_verified AND ai.is_exclusive""",
                             (f"{SN_IDENT_TYPE}.{payload.product_id}", SN_IDENT_TYPE, sns))
            by_sn = {row["sn"]: row["status_code"] for row in rows}
            for sn in sns:
                status = by_sn.get(sn)
                items.append({"serial_number": sn, "exists": status is not None, "status": status})
    return {"items": items}


@router.post("/import-file")
async def import_serial_file(file: UploadFile = File(...), request: Request = None,
                             user: dict[str, Any] = Depends(require_user)) -> dict[str, Any]:
    """从 xlsx / csv 第一列提取 SN 列表（纯解析，不写库），供前端填充登记框。"""
    _csrf(request)
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail=f"导入文件不能超过 {MAX_IMPORT_BYTES // (1024 * 1024)}MB")
    items = _extract_serials_from_file(data, file.filename or "")
    return {"items": items, "count": len(items)}
