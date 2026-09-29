"""CSV / XLSX 导出工具。

CSV 带 utf-8-sig BOM，Excel 打开中文不乱码；XLSX 用 openpyxl（已在依赖中）。
columns 统一为 [(表头, extractor(row))]，让列表渲染与导出列定义完全一致。
"""

from __future__ import annotations

import csv
import io
from typing import Any, Callable
from urllib.parse import quote

from fastapi import HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Font

from .db import fetch_all
from .list_params import parse_ids

MAX_EXPORT_ROWS = 50_000

Column = tuple[str, Callable[[dict[str, Any]], Any]]

# Excel/CSV 公式注入：以这些字符开头的字符串单元格在 Excel 打开时会被当公式执行
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):  # date/datetime → 字符串
        return value.isoformat(sep=" ") if value.__class__.__name__ == "datetime" else value.isoformat()
    if isinstance(value, str) and value[:1] in _FORMULA_PREFIXES:
        # 前置单引号：CSV 里是 Excel 的文本标记（不显示），XLSX 里以普通字符串落库
        return "'" + value
    return value


def csv_bytes(rows: list[dict[str, Any]], columns: list[Column]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([header for header, _ in columns])
    for row in rows:
        writer.writerow([_cell(extractor(row)) for _, extractor in columns])
    # utf-8-sig BOM 兼容 Excel
    return ("﻿" + output.getvalue()).encode("utf-8")


def xlsx_bytes(rows: list[dict[str, Any]], columns: list[Column]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "导出"
    for col_index, (header, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, start=2):
        for col_index, (_, extractor) in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=col_index, value=_cell(extractor(row)))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def attachment_disposition(filename: str) -> str:
    """构造安全 Content-Disposition：ASCII filename + RFC5987 filename*。

    响应头按 latin-1 编码，原始用户文件名（中文/引号/CR/LF）不能直接进
    filename=——统一在这里兜底，导出与附件下载共用。
    """
    filename_enc = quote(filename or "attachment", safe="")
    ascii_name = (filename or "attachment").encode("ascii", "replace").decode("ascii").replace("?", "_")
    ascii_name = ascii_name.replace('"', "_").replace("\r", "").replace("\n", "") or "attachment"
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{filename_enc}'


def export_response(
    rows: list[dict[str, Any]],
    columns: list[Column],
    filename: str,
    fmt: str,
) -> Response:
    """统一导出响应：fmt ∈ {csv, xlsx}；Content-Disposition 含 UTF-8 文件名。"""
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(status_code=422, detail=f"导出行数超过上限 {MAX_EXPORT_ROWS}")
    if fmt == "xlsx":
        content = xlsx_bytes(rows, columns)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ext = "xlsx"
    elif fmt == "csv":
        content = csv_bytes(rows, columns)
        media_type = "text/csv; charset=utf-8"
        ext = "csv"
    else:
        raise HTTPException(status_code=422, detail="fmt 仅支持 csv 或 xlsx")
    return Response(content=content, media_type=media_type,
                    headers={"Content-Disposition": attachment_disposition(f"{filename}.{ext}")})


def export_rows_by_ids(
    conn: Any,
    ids: str,
    id_column: str,
    base_query: str,
    default_order: str,
    columns: list[Column],
    filename: str,
    fmt: str,
) -> Response:
    """按 ids 导出（客户端列表「全部页/多选行」通用）。ids 为空 → 全部行。

    base_query 含 {where} 与 {order_by} 占位符；传入了 ids 时按 ids 顺序
    （array_position）导出，保证与用户当前看到的顺序一致。
    """
    if ids:
        try:
            id_list = parse_ids(ids)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from None
        where = f"{id_column} = ANY(%s)"
        # 保序：与用户看到的勾选顺序一致（psycopg2 把 list 适配为数组；::bigint[] 对齐列类型）
        order_by = f"array_position(%s::bigint[], {id_column})"
        params: list[Any] = [id_list, id_list]
    else:
        where, params, order_by = "TRUE", [], default_order
        # 全量导出在 SQL 层限行（上限+1），超限由 export_response 统一 422，避免先物化
        limit_clause = f" LIMIT {MAX_EXPORT_ROWS + 1}"
        rows = fetch_all(conn, base_query.format(where=where, order_by=order_by) + limit_clause, tuple(params))
        return export_response(rows, columns, filename, fmt)
    rows = fetch_all(conn, base_query.format(where=where, order_by=order_by), tuple(params))
    return export_response(rows, columns, filename, fmt)
